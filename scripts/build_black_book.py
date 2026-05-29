from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Advanced_WMS_Black_Book_With_Database_Tables.docx"
ERD_IMAGE = Path(r"C:\Users\shriyal\Downloads\Warehouse ERD.jpeg.png")
DOC_SKILL = Path(
    r"C:\Users\shriyal\.codex\plugins\cache\openai-primary-runtime\documents\26.430.10722\skills\documents"
)
sys.path.insert(0, str(DOC_SKILL / "scripts"))

from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa


ACCENT = RGBColor(31, 78, 121)
LIGHT_FILL = "EAF2F8"
HEADER_FILL = "D9EAF7"
MUTED = RGBColor(90, 90, 90)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, size=10, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, weights, font_size=9):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        set_cell_text(hdr.cells[idx], header, bold=True, size=font_size, color=RGBColor(0, 0, 0), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[idx], HEADER_FILL)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=font_size)
    widths = column_widths_from_weights(weights, section_content_width_dxa(document.sections[-1]))
    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)
    document.add_paragraph()
    return table


def add_schema_table(document, title, rows):
    document.add_heading(title, level=3)
    add_table(
        document,
        ["Field Name", "Data Type", "Index", "Description"],
        rows,
        [1.5, 1.1, 1.4, 3.0],
        font_size=8,
    )


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(document, items):
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_labeled_paragraph(document, label, text):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(label + ": ")
    run.bold = True
    run.font.color.rgb = ACCENT
    p.add_run(text)


def add_code_block(document, text):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [("Title", 22), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.font.color.rgb = ACCENT

    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = "Advanced Warehouse Management System - Black Book"
    header_para.runs[0].font.size = Pt(9)
    header_para.runs[0].font.color.rgb = MUTED

    footer = section.footer
    add_page_number(footer.paragraphs[0])
    return doc


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Advanced Warehouse Management System")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = ACCENT

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Black Book Project Documentation")
    r.font.size = Pt(15)

    doc.add_paragraph()
    meta_rows = [
        ("Project Type", "Web Application"),
        ("Domain", "Logistics, Warehouse, Inventory and Supply Chain Management"),
        ("Backend", "Java 21+, Spring Boot 3, Spring Data JPA, Spring Security JWT"),
        ("Frontend", "React.js with Vite"),
        ("Database", "PostgreSQL"),
        ("Prepared For", "Infotact Technical Internship Program"),
    ]
    add_table(doc, ["Field", "Details"], meta_rows, [1.2, 3.8], font_size=10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Prepared by: Shriyal").bold = True
    doc.add_page_break()


def add_requirement_summary(doc):
    doc.add_heading("Project Requirement Summary", level=1)
    rows = [
        ("Technology Stack", "React.js frontend, Java Spring Boot backend, REST API, JWT security, Docker support and GitHub Actions CI."),
        ("Database", "PostgreSQL relational database with JPA entities, relationships, constraints and transaction-safe inventory updates."),
        ("Testing", "JUnit 5 and Spring Boot integration tests for warehouse, receiving and order workflows; frontend production build verification; user acceptance flow through the React UI."),
    ]
    add_table(doc, ["Project Requirement", "Details"], rows, [1.4, 3.6], font_size=10)
    doc.add_paragraph(
        "This report follows the Development index provided for web/mobile application documentation. "
        "The content is customized for the Advanced Warehouse Management System implemented in the current project repository."
    )
    doc.add_paragraph(
        "The purpose of this black book is to explain the project in a simple and complete way. It covers why the "
        "project is needed, what problem it solves, how the design is prepared, how the frontend and backend are "
        "implemented, how the database stores information, and how testing is done before submission."
    )
    doc.add_page_break()


def add_index(doc):
    doc.add_heading("Index: Development (Web/Mobile Applications, ETL)", level=1)
    rows = [
        ("Chapter 1", "Introduction", ""),
        ("1.1", "Problem Statement", ""),
        ("1.2", "Objectives", ""),
        ("1.3", "Scope", ""),
        ("Chapter 2", "Design", ""),
        ("2.1", "System Architecture", ""),
        ("2.2", "Database Design", ""),
        ("Chapter 3", "Implementation", ""),
        ("3.1", "Frontend Development", ""),
        ("3.2", "Backend Development", ""),
        ("3.3", "Integration", ""),
        ("Chapter 4", "Testing", ""),
        ("4.1", "Test Cases", ""),
        ("4.2", "Results", ""),
        ("Chapter 5", "Conclusion", ""),
        ("5.1", "Summary", ""),
        ("5.2", "Future Enhancements", ""),
        ("Chapter 6", "References", ""),
        ("Chapter 7", "Appendices", ""),
        ("Chapter 8", "Annexure - Progress Sheet", ""),
    ]
    add_table(doc, ["Chapter", "Content", "Page Number"], rows, [1.1, 3.2, 1.2], font_size=10)
    doc.add_page_break()


def chapter_1(doc):
    doc.add_heading("Chapter 1: Introduction", level=1)
    doc.add_paragraph(
        "The Advanced Warehouse Management System is a web-based enterprise application designed to automate "
        "warehouse operations such as product catalog management, storage hierarchy maintenance, inventory receiving, "
        "putaway, stock tracking and customer order fulfillment. The project is based on the Enterprise Warehouse "
        "Management System brief from the Infotact Technical Internship Program."
    )
    doc.add_paragraph(
        "Modern warehouses handle large volumes of products, multiple storage locations and fast order turnaround. "
        "Manual spreadsheet-based tracking can cause stock mismatch, misplaced items, slow picking and delayed "
        "dispatch. This system solves those problems by connecting a React dashboard with secure Spring Boot REST APIs "
        "and a PostgreSQL relational database."
    )
    add_labeled_paragraph(
        doc,
        "Meaning of WMS",
        "A Warehouse Management System is software that helps a company manage goods inside a warehouse. It keeps "
        "records of where products are stored, how much stock is available, which orders are pending, and which "
        "items need to be picked, packed or shipped."
    )
    add_labeled_paragraph(
        doc,
        "Need of the project",
        "In a manual warehouse, workers may write stock details in registers or spreadsheets. This is slow and can "
        "lead to mistakes. A web-based system updates information immediately and gives all users the same latest "
        "data. This improves speed, accuracy and control."
    )
    add_labeled_paragraph(
        doc,
        "Proposed solution",
        "The proposed system provides a login-based dashboard where an administrator or operator can manage products, "
        "warehouse locations, incoming stock and customer orders. The backend applies business rules so that stock "
        "changes happen safely and the database remains consistent."
    )

    doc.add_heading("1.1 Problem Statement", level=2)
    doc.add_paragraph(
        "Traditional warehouse operations often depend on manual registers or disconnected spreadsheets. This creates "
        "operational bottlenecks: inventory is not updated in real time, stock can be stored in the wrong bin, and "
        "orders may be accepted even when stock is unavailable. These issues lead to stockouts, delayed order packing "
        "and poor visibility for warehouse managers."
    )
    doc.add_paragraph(
        "The problem is to design and implement a secure, database-driven WMS that maintains accurate inventory data, "
        "assigns inbound stock to available storage bins, supports barcode/QR based identification and guides orders "
        "through a controlled fulfillment workflow."
    )
    doc.add_paragraph(
        "The main challenge in warehouse software is not only storing product details. The system must also know the "
        "correct quantity, location and status of every item. If two users update the same stock at the same time, "
        "the final quantity should still be correct. For this reason, the project uses transactions and database "
        "locking during important inventory operations."
    )
    doc.add_paragraph(
        "Another challenge is easy use by warehouse staff. Floor operators need clear screens and simple forms because "
        "warehouse work is time-sensitive. The React frontend is therefore designed as a practical dashboard with "
        "separate tabs for common tasks."
    )

    doc.add_heading("1.2 Objectives", level=2)
    add_bullets(doc, [
        "Create a hierarchical warehouse model with warehouses, zones, aisles and storage bins.",
        "Maintain a product catalog with SKU, barcode, category, price, weight and threshold details.",
        "Implement transactional receiving and putaway logic that updates stock and bin capacity atomically.",
        "Maintain inventory snapshots and ledger-style inventory transactions.",
        "Provide an order workflow from pending to picking, packed and shipped.",
        "Generate QR code images for product identification.",
        "Secure APIs using JWT authentication and role-based access for ADMIN and OPERATOR users.",
        "Build a React dashboard for warehouse staff and managers.",
        "Use PostgreSQL for relational data integrity and Spring Data JPA for data access.",
        "Verify the system through automated backend tests and frontend build validation.",
    ])
    doc.add_paragraph(
        "These objectives make the system useful for both management-level users and floor-level users. A manager "
        "gets visibility of warehouse capacity and stock, while an operator gets a guided workflow for receiving, "
        "picking and order processing."
    )
    add_labeled_paragraph(
        doc,
        "Functional objective",
        "The system should perform real warehouse tasks such as adding products, receiving stock, selecting storage "
        "bins, showing inventory and processing orders."
    )
    add_labeled_paragraph(
        doc,
        "Technical objective",
        "The system should be built using modern Java web development practices such as REST APIs, JPA, validation, "
        "JWT security, PostgreSQL and automated tests."
    )
    add_labeled_paragraph(
        doc,
        "Quality objective",
        "The system should be easy to understand, easy to run locally, and organized in a way that future developers "
        "can maintain and extend."
    )

    doc.add_heading("1.3 Scope", level=2)
    doc.add_paragraph("The project scope includes the following modules:")
    add_bullets(doc, [
        "Authentication and role-based authorization.",
        "Warehouse master data management.",
        "Zone, aisle and storage bin management.",
        "Product and product category management.",
        "Receiving, putaway and inventory snapshot viewing.",
        "Customer order creation and fulfillment state transitions.",
        "Supplier and purchase order management.",
        "User management for administrators.",
        "Docker and local development execution support.",
    ])
    doc.add_paragraph(
        "The current scope does not include live handheld scanner hardware integration, advanced route optimization, "
        "machine-learning demand forecasting or payment/accounting modules. These are listed as future enhancements."
    )
    add_labeled_paragraph(
        doc,
        "Users in scope",
        "The main users are ADMIN and OPERATOR. ADMIN users manage master data, users and procurement. OPERATOR users "
        "mainly work with receiving, inventory and order movement."
    )
    add_labeled_paragraph(
        doc,
        "Data in scope",
        "The system stores warehouse locations, product details, inventory quantities, stock movements, customer orders, "
        "purchase orders, suppliers and application users."
    )
    add_labeled_paragraph(
        doc,
        "Security in scope",
        "Only logged-in users can access protected APIs. Admin-only features are restricted with role checks. Passwords "
        "are stored as hashes and JWT tokens are used for secure API communication."
    )
    doc.add_page_break()


def chapter_2(doc):
    doc.add_heading("Chapter 2: Design", level=1)
    doc.add_paragraph(
        "The system is designed as a layered web application. The frontend is responsible for user interaction and "
        "form validation, while the backend handles business rules, security, persistence and transactional integrity. "
        "PostgreSQL acts as the central source of truth for all warehouse, product, inventory and order data."
    )
    doc.add_paragraph(
        "A layered design is used because it keeps the project clean. The user interface, business logic and database "
        "logic are separated from each other. This means that if the frontend design changes, the database code does "
        "not need to change. Similarly, if a database query changes, the user screens can remain mostly the same."
    )
    add_labeled_paragraph(
        doc,
        "Design approach",
        "The application follows a common web architecture: React runs in the browser, Spring Boot runs on the server, "
        "and PostgreSQL stores the data. The browser never talks directly to the database. It always talks to the "
        "backend API, which checks security and applies rules."
    )

    doc.add_heading("2.1 System Architecture", level=2)
    arch_rows = [
        ("Presentation Layer", "React.js, Vite, CSS, Lucide icons", "Displays dashboards, forms and tables for users."),
        ("API Layer", "Spring Boot REST Controllers", "Exposes endpoints for authentication, products, warehouses, inventory, orders, procurement and users."),
        ("Service Layer", "Spring Services with @Transactional", "Implements receiving, putaway, order packing, user management and procurement workflows."),
        ("Data Access Layer", "Spring Data JPA repositories", "Executes parameterized queries and entity persistence through Hibernate."),
        ("Database Layer", "PostgreSQL", "Stores relational data with foreign keys, constraints and transaction support."),
        ("Security Layer", "Spring Security, JWT, BCrypt", "Authenticates users and restricts actions by ADMIN and OPERATOR roles."),
    ]
    add_table(doc, ["Layer", "Technology", "Responsibility"], arch_rows, [1.3, 1.7, 3.0], font_size=9)
    doc.add_paragraph(
        "The architecture makes the system easier to test and extend. For example, if a new report module is needed, "
        "a new React view and backend endpoint can be added without rewriting the whole application. The same database "
        "entities and repositories can be reused."
    )

    doc.add_paragraph("High level request flow:")
    add_numbered(doc, [
        "User logs in from the React frontend.",
        "Spring Security validates credentials and returns a JWT token.",
        "Frontend sends the token in the Authorization header for protected API calls.",
        "Controller validates the request and forwards work to the service layer.",
        "Service layer applies business rules inside a transaction.",
        "Repository layer reads/writes PostgreSQL tables through JPA.",
        "Backend returns JSON responses consumed by the React dashboard.",
    ])
    tech_rows = [
        ("React.js", "Used to build a dynamic frontend where screens update without full page reloads."),
        ("Vite", "Used as a fast development server and build tool for the React frontend."),
        ("Spring Boot", "Used to create production-ready REST APIs quickly with built-in server support."),
        ("Spring Data JPA", "Used to map Java classes to database tables and reduce manual SQL code."),
        ("Spring Security", "Used for login, JWT validation and role-based authorization."),
        ("PostgreSQL", "Used because it is reliable for relational data and supports transactions and constraints."),
        ("JUnit 5", "Used to test important backend services and protect the project from regressions."),
        ("Docker Compose", "Used to run the application with PostgreSQL in a repeatable environment."),
    ]
    add_table(doc, ["Technology", "Why It Is Used"], tech_rows, [1.5, 4.5], font_size=9)
    add_labeled_paragraph(
        doc,
        "Client-server communication",
        "The React application sends HTTP requests to Spring Boot endpoints. The backend returns JSON data. This format "
        "is easy for the frontend to read and display in tables, forms and dashboards."
    )
    add_labeled_paragraph(
        doc,
        "Security flow",
        "After login, the backend returns a JWT token. The frontend sends this token with every protected request. "
        "Spring Security reads the token and decides whether the user is allowed to perform the requested action."
    )

    doc.add_heading("2.2 Database Design", level=2)
    doc.add_paragraph(
        "The database is normalized around warehouse location, product catalog, inventory, procurement, users and "
        "customer orders. The key relationship is the physical storage hierarchy: one warehouse contains many zones, "
        "one zone contains many aisles, and one aisle contains many storage bins."
    )
    doc.add_paragraph(
        "Database design is important in this project because inventory data must be accurate. If the database allows "
        "wrong or duplicate data, the warehouse staff may pick the wrong item or accept an order without enough stock. "
        "For this reason, the design uses primary keys, foreign keys, unique fields and transaction-safe updates."
    )
    entity_rows = [
        ("warehouses", "Stores warehouse code, name, address/location and active status."),
        ("zones", "Groups areas inside a warehouse and links to warehouses."),
        ("aisles", "Represents aisle-level storage organization inside a zone."),
        ("storage_bins", "Stores bin code, capacity, current occupancy, active flag and status."),
        ("product_categories", "Classifies products and can define a preferred storage zone."),
        ("products", "Stores SKU, barcode, name, description, price, weight, threshold and category."),
        ("inventory_items", "Stores product quantity by storage bin, batch number, expiry date and reserved quantity."),
        ("inventory_transactions", "Logs RECEIVE and PICK movements for auditability."),
        ("customer_orders", "Stores customer order number, warehouse and fulfillment status."),
        ("customer_order_lines", "Stores ordered products and requested quantities."),
        ("suppliers", "Stores supplier contact and address data."),
        ("purchase_orders", "Stores supplier purchase order status and expected date."),
        ("purchase_order_items", "Stores products and quantities ordered from suppliers."),
        ("app_users", "Stores application users, roles, status, contact details and password hashes."),
    ]
    add_table(doc, ["Table", "Purpose"], entity_rows, [1.5, 4.5], font_size=9)
    doc.add_paragraph(
        "The inventory_items table stores the current stock position, while inventory_transactions stores the history "
        "of stock movement. This is useful because the system can show both the present quantity and the reason behind "
        "past changes. For example, a RECEIVE transaction increases stock and a PICK transaction decreases stock."
    )

    rel_rows = [
        ("Warehouse -> Zone", "One-to-many", "A warehouse can contain many zones."),
        ("Zone -> Aisle", "One-to-many", "A zone can contain many aisles."),
        ("Aisle -> StorageBin", "One-to-many", "An aisle can contain many storage bins."),
        ("ProductCategory -> Product", "One-to-many", "A category can classify many products."),
        ("Product + StorageBin -> InventoryItem", "Many-to-one links", "Inventory is tracked for each product stored in each bin."),
        ("CustomerOrder -> CustomerOrderLine", "One-to-many", "An order can contain multiple product lines."),
        ("PurchaseOrder -> PurchaseOrderItem", "One-to-many", "A purchase order can contain multiple product lines."),
    ]
    add_table(doc, ["Relationship", "Type", "Description"], rel_rows, [1.8, 1.2, 3.0], font_size=9)
    concept_rows = [
        ("Primary Key", "A unique value that identifies one row in a table, such as an id."),
        ("Foreign Key", "A reference from one table to another, such as product_id in inventory_items."),
        ("Normalization", "Dividing data into related tables to avoid repeated and inconsistent data."),
        ("Transaction", "A group of database operations that must succeed or fail together."),
        ("ACID", "A database reliability concept that helps keep inventory data correct."),
        ("Constraint", "A database rule that prevents invalid data, such as duplicate SKU values."),
    ]
    add_table(doc, ["Database Concept", "Simple Meaning in This Project"], concept_rows, [1.5, 4.5], font_size=9)
    add_labeled_paragraph(
        doc,
        "Why PostgreSQL",
        "PostgreSQL is suitable for this system because warehouse data is strongly relational. Products are linked "
        "to categories, inventory is linked to products and bins, and orders are linked to order lines. PostgreSQL "
        "handles these relationships safely."
    )
    add_labeled_paragraph(
        doc,
        "Why JPA/Hibernate",
        "JPA and Hibernate allow developers to work with Java objects instead of writing SQL for every operation. "
        "This keeps the backend code cleaner and helps prevent common SQL mistakes."
    )

    if ERD_IMAGE.exists():
        doc.add_paragraph("Figure 1: Warehouse Management System ERD")
        picture = doc.add_picture(str(ERD_IMAGE), width=Inches(6.3))
        picture._inline.docPr.set("descr", "Entity relationship diagram for the Advanced Warehouse Management System database.")
        picture._inline.docPr.set("title", "Advanced WMS ERD")
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Detailed Database Table Structure", level=3)
    doc.add_paragraph(
        "The following table definitions show the database design in field-wise format. Each table includes the field "
        "name, expected data type, key/index information and a simple description. The format is kept similar to the "
        "college black book table style."
    )
    schema_tables = [
        (
            "2.2.1 Warehouse Table (warehouses)",
            [
                ("id", "BIGINT", "Primary Key", "Unique warehouse record id."),
                ("code", "VARCHAR", "Not Null, Unique", "Short warehouse code such as BLR-01."),
                ("name", "VARCHAR", "Not Null", "Warehouse name."),
                ("address", "VARCHAR", "Nullable", "Warehouse address or location detail."),
            ],
        ),
        (
            "2.2.2 Zone Table (zones)",
            [
                ("id", "BIGINT", "Primary Key", "Unique zone record id."),
                ("code", "VARCHAR", "Not Null", "Zone code used inside a warehouse."),
                ("name", "VARCHAR", "Not Null", "Readable zone name."),
                ("warehouse_id", "BIGINT", "Foreign Key, Not Null", "References the warehouse where the zone belongs."),
            ],
        ),
        (
            "2.2.3 Aisle Table (aisles)",
            [
                ("id", "BIGINT", "Primary Key", "Unique aisle record id."),
                ("code", "VARCHAR", "Not Null", "Aisle code used for physical storage location."),
                ("zone_id", "BIGINT", "Foreign Key, Not Null", "References the zone where the aisle belongs."),
            ],
        ),
        (
            "2.2.4 Storage Bin Table (storage_bins)",
            [
                ("id", "BIGINT", "Primary Key", "Unique storage bin record id."),
                ("code", "VARCHAR", "Not Null, Unique", "Bin code printed or displayed for storage location."),
                ("capacity", "INTEGER", "Not Null", "Total capacity units of the bin."),
                ("used_capacity", "INTEGER", "Not Null", "Currently occupied capacity units."),
                ("active", "BOOLEAN", "Default True", "Shows whether the bin is active for use."),
                ("status", "VARCHAR", "Enum", "Bin status such as AVAILABLE, FULL or INACTIVE."),
                ("version", "BIGINT", "Version", "Used for optimistic locking."),
                ("aisle_id", "BIGINT", "Foreign Key, Not Null", "References the aisle where the bin is located."),
            ],
        ),
        (
            "2.2.5 Product Category Table (product_categories)",
            [
                ("id", "BIGINT", "Primary Key", "Unique product category id."),
                ("created_at", "TIMESTAMP", "Not Null", "Date and time when category was created."),
                ("updated_at", "TIMESTAMP", "Not Null", "Date and time when category was last updated."),
                ("active", "BOOLEAN", "Default True", "Shows whether the category is active."),
                ("name", "VARCHAR", "Not Null", "Category name such as Electronics or Grocery."),
                ("preferred_zone_id", "BIGINT", "Foreign Key", "Preferred zone for products in this category."),
                ("parent_category_id", "BIGINT", "Foreign Key", "Parent category for category hierarchy."),
                ("warehouse_id", "BIGINT", "Foreign Key, Not Null", "Warehouse where this category is used."),
            ],
        ),
        (
            "2.2.6 Product Table (products)",
            [
                ("id", "BIGINT", "Primary Key", "Unique product id."),
                ("sku", "VARCHAR", "Not Null, Unique, Index", "Stock keeping unit used to identify product."),
                ("name", "VARCHAR", "Not Null", "Product name."),
                ("barcode", "VARCHAR", "Unique, Index", "Barcode or QR value for product scanning."),
                ("description", "VARCHAR", "Nullable", "Short product description."),
                ("unit_volume", "INTEGER", "Not Null", "Capacity units consumed by one product item."),
                ("min_threshold", "NUMERIC", "Not Null", "Minimum stock level or reorder threshold."),
                ("price", "NUMERIC", "Default 0", "Product price."),
                ("weight", "DOUBLE", "Nullable", "Product weight."),
                ("active", "BOOLEAN", "Default True", "Shows whether product is active."),
                ("category_id", "BIGINT", "Foreign Key", "References product category."),
                ("warehouse_id", "BIGINT", "Foreign Key", "References default warehouse."),
            ],
        ),
        (
            "2.2.7 Inventory Item Table (inventory_items)",
            [
                ("id", "BIGINT", "Primary Key", "Unique inventory item id."),
                ("product_id", "BIGINT", "Foreign Key, Not Null", "References stored product."),
                ("storage_bin_id", "BIGINT", "Foreign Key, Not Null", "References bin where stock is stored."),
                ("quantity", "INTEGER", "Not Null", "Current quantity in the bin."),
                ("reserved_quantity", "INTEGER", "Nullable", "Quantity reserved for orders."),
                ("batch_number", "VARCHAR", "Nullable", "Batch number of received stock."),
                ("expiry_date", "DATE", "Nullable", "Expiry date for perishable stock."),
                ("status", "VARCHAR", "Enum", "Inventory status such as AVAILABLE."),
                ("created_at", "TIMESTAMP", "Nullable", "Date and time when stock row was created."),
                ("updated_at", "TIMESTAMP", "Nullable", "Date and time when stock row was updated."),
                ("version", "BIGINT", "Version", "Used for optimistic locking."),
            ],
        ),
        (
            "2.2.8 Inventory Transaction Table (inventory_transactions)",
            [
                ("id", "BIGINT", "Primary Key", "Unique transaction id."),
                ("type", "VARCHAR", "Not Null, Enum", "Transaction type such as RECEIVE or PICK."),
                ("product_id", "BIGINT", "Foreign Key, Not Null", "Product affected by transaction."),
                ("storage_bin_id", "BIGINT", "Foreign Key, Not Null", "Storage bin affected by transaction."),
                ("quantity_delta", "INTEGER", "Not Null", "Positive or negative stock movement."),
                ("reference", "VARCHAR", "Nullable", "Reference such as ASN or order number."),
                ("created_at", "TIMESTAMP", "Not Null", "Date and time of transaction."),
            ],
        ),
        (
            "2.2.9 Customer Order Table (customer_orders)",
            [
                ("id", "BIGINT", "Primary Key", "Unique order id."),
                ("order_number", "VARCHAR", "Not Null, Unique", "Human-readable customer order number."),
                ("status", "VARCHAR", "Not Null, Enum", "Order status such as PENDING, PICKING, PACKED or SHIPPED."),
                ("created_at", "TIMESTAMP", "Not Null", "Date and time when order was created."),
                ("expected_ship_date", "TIMESTAMP", "Nullable", "Expected date for shipment."),
                ("packed_at", "TIMESTAMP", "Nullable", "Date and time when order was packed."),
                ("shipped_at", "TIMESTAMP", "Nullable", "Date and time when order was shipped."),
                ("warehouse_id", "BIGINT", "Foreign Key", "Warehouse from which order is fulfilled."),
            ],
        ),
        (
            "2.2.10 Customer Order Line Table (customer_order_lines)",
            [
                ("id", "BIGINT", "Primary Key", "Unique order line id."),
                ("customer_order_id", "BIGINT", "Foreign Key, Not Null", "References the customer order."),
                ("product_id", "BIGINT", "Foreign Key, Not Null", "References ordered product."),
                ("requested_quantity", "INTEGER", "Not Null", "Quantity requested by customer."),
                ("picked_quantity", "INTEGER", "Not Null", "Quantity picked during fulfillment."),
            ],
        ),
        (
            "2.2.11 Supplier Table (suppliers)",
            [
                ("id", "BIGINT", "Primary Key", "Unique supplier id."),
                ("created_at", "TIMESTAMP", "Not Null", "Date and time when supplier was created."),
                ("updated_at", "TIMESTAMP", "Not Null", "Date and time when supplier was updated."),
                ("active", "BOOLEAN", "Default True", "Shows whether supplier is active."),
                ("name", "VARCHAR", "Not Null, Unique", "Supplier name."),
                ("address", "VARCHAR", "Nullable", "Supplier address."),
                ("contact_email", "VARCHAR", "Nullable", "Supplier email address."),
                ("phone", "VARCHAR", "Nullable", "Supplier contact number."),
            ],
        ),
        (
            "2.2.12 Purchase Order Table (purchase_orders)",
            [
                ("id", "BIGINT", "Primary Key", "Unique purchase order id."),
                ("order_date", "TIMESTAMP", "Not Null", "Date when purchase order was created."),
                ("expected_date", "TIMESTAMP", "Nullable", "Expected delivery date."),
                ("status", "VARCHAR", "Not Null, Enum", "Purchase order status such as ORDERED or RECEIVED."),
                ("supplier_id", "BIGINT", "Foreign Key, Not Null", "References supplier."),
                ("warehouse_id", "BIGINT", "Foreign Key, Not Null", "Warehouse receiving the purchase order."),
            ],
        ),
        (
            "2.2.13 Purchase Order Item Table (purchase_order_items)",
            [
                ("id", "BIGINT", "Primary Key", "Unique purchase order item id."),
                ("purchase_order_id", "BIGINT", "Foreign Key, Not Null", "References purchase order."),
                ("product_id", "BIGINT", "Foreign Key, Not Null", "References product being purchased."),
                ("quantity", "INTEGER", "Not Null", "Quantity ordered from supplier."),
            ],
        ),
        (
            "2.2.14 Application User Table (app_users)",
            [
                ("id", "BIGINT", "Primary Key", "Unique user id."),
                ("username", "VARCHAR", "Not Null, Unique", "Login username."),
                ("created_at", "TIMESTAMP", "Nullable", "Date and time when user was created."),
                ("updated_at", "TIMESTAMP", "Nullable", "Date and time when user was updated."),
                ("name", "VARCHAR", "Nullable", "Full name of user."),
                ("email", "VARCHAR", "Unique", "User email address."),
                ("contact_number", "VARCHAR", "Nullable", "User phone number."),
                ("password_hash", "VARCHAR", "Not Null", "BCrypt hashed password."),
                ("role", "VARCHAR", "Not Null, Enum", "User role such as ADMIN or OPERATOR."),
                ("status", "VARCHAR", "Default ACTIVE", "User account status."),
                ("warehouse_id", "BIGINT", "Foreign Key", "Warehouse assigned to user."),
            ],
        ),
    ]
    for title, rows in schema_tables:
        add_schema_table(doc, title, rows)
    doc.add_page_break()


def chapter_3(doc):
    doc.add_heading("Chapter 3: Implementation", level=1)
    doc.add_paragraph(
        "Implementation is divided into frontend, backend and integration work. The codebase contains a Spring Boot "
        "backend under the backend folder and a React/Vite frontend under the frontend folder."
    )
    doc.add_paragraph(
        "The implementation is done in a practical way so that each module can be understood separately. The frontend "
        "focuses on screens and user actions. The backend focuses on validation, security and business rules. The "
        "database stores the final data permanently."
    )

    doc.add_heading("3.1 Frontend Development", level=2)
    doc.add_paragraph(
        "The frontend is implemented using React.js and Vite. It provides a dashboard-style interface for warehouse "
        "operators and administrators. The UI is organized into modules such as login, dashboard, warehouses, products, "
        "receiving, inventory, orders, procurement and users."
    )
    doc.add_paragraph(
        "React was selected because it is component-based. A component is a reusable part of the screen, such as a "
        "form, table, tab or button. This makes the interface easier to maintain because the same style and behavior "
        "can be reused in different places."
    )
    frontend_rows = [
        ("Login", "Collects username/password and stores the JWT token for API calls."),
        ("Dashboard", "Shows high-level operational information for quick monitoring."),
        ("Warehouses", "Allows viewing and managing warehouses, zones, aisles and bins."),
        ("Products", "Allows product creation and displays barcode/QR information."),
        ("Receiving", "Receives stock into the system and triggers putaway logic."),
        ("Inventory", "Displays stock snapshots by product, bin and warehouse."),
        ("Orders", "Creates orders and moves them through picking, packing and shipping."),
        ("Procurement", "Manages suppliers and purchase orders."),
        ("Users", "Allows ADMIN users to create and edit users."),
    ]
    add_table(doc, ["Frontend Module", "Purpose"], frontend_rows, [1.4, 4.6], font_size=9)
    doc.add_paragraph(
        "Form-level validation is used to reduce invalid submissions. Examples include required fields, positive "
        "numeric quantities, SKU/barcode format patterns, email format validation, password length rules and role/status selection."
    )
    doc.add_paragraph(
        "The frontend stores the login token after a successful login. When the user opens protected screens, the token "
        "is sent to the backend in the Authorization header. If the token is missing or invalid, the backend refuses "
        "the request. This keeps warehouse data protected from unauthorized access."
    )
    frontend_validation_rows = [
        ("Required fields", "Important fields like username, SKU, product name and quantity cannot be blank."),
        ("Positive numbers", "Quantity, capacity, price and weight must be valid positive values."),
        ("Pattern checks", "SKU and barcode values use safe characters to avoid invalid input."),
        ("Role checks", "Admin-only tabs such as users and procurement are not shown to normal operators."),
        ("Clear feedback", "Errors returned by the backend are shown to the user so the issue can be corrected."),
    ]
    add_table(doc, ["Frontend Validation", "Simple Explanation"], frontend_validation_rows, [1.7, 4.3], font_size=9)
    add_labeled_paragraph(
        doc,
        "User experience",
        "The interface is designed for regular warehouse work. The main modules are available as tabs, and forms are "
        "kept simple so users can quickly complete common operations like receiving stock or processing orders."
    )

    doc.add_heading("3.2 Backend Development", level=2)
    doc.add_paragraph(
        "The backend is a Spring Boot 3 application. It uses controllers for REST endpoints, DTO records for request "
        "and response data, services for business logic, repositories for persistence and entities for database mapping."
    )
    doc.add_paragraph(
        "Spring Boot reduces boilerplate code and gives the application an embedded server. Because of this, the backend "
        "can be started easily from the command line. The backend also provides a clear package structure so that "
        "controllers, services, repositories, domain classes and security classes are easy to find."
    )
    backend_rows = [
        ("AuthController", "/api/auth", "Login and JWT token generation."),
        ("WarehouseController", "/api/warehouses, /api/zones, /api/aisles, /api/bins", "Warehouse hierarchy CRUD."),
        ("ProductController", "/api/products", "Product catalog and QR/barcode generation."),
        ("ProductCategoryController", "/api/product-categories", "Product category management."),
        ("InventoryController", "/api/inventory", "Receiving and inventory snapshot APIs."),
        ("OrderController", "/api/orders", "Order creation and fulfillment workflow."),
        ("ProcurementController", "/api/procurement", "Supplier and purchase order workflow."),
        ("UserController", "/api/users", "Admin-only user management."),
    ]
    add_table(doc, ["Controller", "Endpoint", "Responsibility"], backend_rows, [1.4, 2.1, 2.5], font_size=8)

    doc.add_paragraph("Important backend implementation points:")
    add_bullets(doc, [
        "Spring Security authenticates users and validates JWT tokens for protected requests.",
        "BCryptPasswordEncoder stores passwords as hashes instead of plain text.",
        "DTO validation uses Jakarta Bean Validation annotations such as @NotBlank, @NotNull and @Positive.",
        "InventoryService uses @Transactional to keep receiving and order picking atomic.",
        "Pessimistic write locks are used while selecting storage bins and inventory rows for stock updates.",
        "GlobalExceptionHandler converts validation, duplicate resource and insufficient stock errors into clean API responses.",
        "Spring Data JPA prevents SQL injection by using repository methods and parameterized JPQL queries.",
        "The application exposes health checks through Spring Boot Actuator.",
    ])
    backend_package_rows = [
        ("domain", "Contains JPA entity classes such as Product, Warehouse, StorageBin and InventoryItem."),
        ("repository", "Contains Spring Data JPA interfaces used to read and write database records."),
        ("service", "Contains business logic such as receiving stock, packing orders and managing users."),
        ("api", "Contains REST controllers that receive frontend requests."),
        ("api.dto", "Contains request and response objects used by the API."),
        ("security", "Contains JWT, password encoder, authentication and role configuration."),
        ("exception", "Contains custom exceptions such as InsufficientStockException."),
        ("config", "Contains application setup such as sample data seeding."),
    ]
    add_table(doc, ["Backend Package", "Purpose"], backend_package_rows, [1.5, 4.5], font_size=9)
    add_labeled_paragraph(
        doc,
        "Receiving logic",
        "When stock is received, the backend first checks the product, finds a storage bin with enough capacity, locks "
        "the required rows, updates inventory, updates bin capacity and records an inventory transaction."
    )
    add_labeled_paragraph(
        doc,
        "Order logic",
        "When an order is packed, the backend checks available inventory. If stock is available, it decreases quantity "
        "and records the movement. If stock is not enough, the operation fails and previous changes are rolled back."
    )
    add_labeled_paragraph(
        doc,
        "Exception handling",
        "Errors are handled in one place using a global exception handler. This gives the frontend clean error messages "
        "instead of raw server errors."
    )

    doc.add_heading("3.3 Integration", level=2)
    doc.add_paragraph(
        "Integration is handled through JSON-based REST APIs. The frontend sends requests to the backend using the "
        "configured base API URL. In local development, the frontend runs on port 5173 and the backend local profile "
        "runs on port 8081."
    )
    doc.add_paragraph(
        "The integration between frontend and backend is important because the React application is only a user "
        "interface. The actual data is stored and processed by the backend. Every create, update or view action from "
        "the frontend becomes an API request to Spring Boot."
    )
    integration_rows = [
        ("Frontend URL", "http://localhost:5173"),
        ("Backend URL", "http://localhost:8081"),
        ("Backend health", "http://localhost:8081/actuator/health"),
        ("Database", "PostgreSQL database wms on localhost:5432"),
        ("API documentation", "Swagger UI available when backend runs on the configured backend port."),
    ]
    add_table(doc, ["Item", "Value"], integration_rows, [1.5, 4.5], font_size=9)
    doc.add_paragraph("Local run commands:")
    add_code_block(doc, "scripts\\run-backend-local.cmd\nscripts\\run-frontend-local.cmd")
    doc.add_paragraph("Typical API flow:")
    add_numbered(doc, [
        "POST /api/auth/login to receive JWT token.",
        "Create or use seeded warehouse master data.",
        "Create products with POST /api/products.",
        "Receive inbound stock with POST /api/inventory/receive.",
        "Create orders with POST /api/orders.",
        "Move orders through /start-picking, /pack and /ship.",
    ])
    integration_theory_rows = [
        ("JSON", "Used as the data format between frontend and backend."),
        ("REST API", "Used to expose backend operations through URLs and HTTP methods."),
        ("CORS", "Allows the React development server to call the Spring Boot backend safely."),
        ("JWT Header", "Sends the logged-in user's token with protected requests."),
        ("Environment Variables", "Used for database URL, database username, password and JWT secret."),
        ("Docker Compose", "Can start PostgreSQL, backend and frontend together for a full-stack run."),
    ]
    add_table(doc, ["Integration Item", "Explanation"], integration_theory_rows, [1.5, 4.5], font_size=9)
    doc.add_paragraph(
        "This integration style is flexible because different frontends can use the same backend later. For example, "
        "a mobile app for floor operators could call the same inventory and order APIs in the future."
    )
    doc.add_page_break()


def chapter_4(doc):
    doc.add_heading("Chapter 4: Testing", level=1)
    doc.add_paragraph(
        "Testing verifies that the system works according to the requirement document. The project uses automated "
        "backend tests, frontend production build validation and manual user acceptance flows through the browser."
    )
    doc.add_paragraph(
        "Testing is very important for a warehouse system because wrong stock data can create real business problems. "
        "If receiving, packing or order status logic fails, the warehouse may show incorrect quantity or ship the wrong "
        "items. For this reason, the project checks both normal cases and error cases."
    )
    testing_type_rows = [
        ("Unit Testing", "Checks a small part of the system logic, such as service behavior or validation."),
        ("Integration Testing", "Checks whether Spring Boot, JPA repositories and the test database work together."),
        ("User Acceptance Testing", "Checks whether the screens support real user tasks such as login, receive stock and pack orders."),
        ("Security Testing", "Checks whether protected APIs require login and admin-only operations are restricted."),
        ("Build Testing", "Checks whether the frontend can be compiled for production without errors."),
        ("Performance Check", "Measures important API response time against the target mentioned in the project brief."),
    ]
    add_table(doc, ["Testing Type", "Simple Explanation"], testing_type_rows, [1.6, 4.4], font_size=9)

    doc.add_heading("4.1 Test Cases", level=2)
    rows = [
        ("TC-01", "Login", "Enter valid admin credentials.", "JWT token is generated and user enters dashboard.", "Pass"),
        ("TC-02", "Warehouse creation", "Create a warehouse with a unique code.", "Warehouse is saved and listed.", "Pass"),
        ("TC-03", "Duplicate warehouse validation", "Create a warehouse with an existing code.", "DuplicateResourceException response is returned.", "Pass"),
        ("TC-04", "Zone creation", "Create a zone under a valid warehouse.", "Zone is linked with the selected warehouse.", "Pass"),
        ("TC-05", "Bin creation", "Create a storage bin under a valid aisle.", "Bin is saved with capacity and status.", "Pass"),
        ("TC-06", "Product creation", "Submit product SKU, name, price and volume.", "Product is saved and QR/barcode can be generated.", "Pass"),
        ("TC-07", "Receiving stock", "Receive product quantity against a reference.", "Stock is placed in an available bin and inventory increases.", "Pass"),
        ("TC-08", "Order packing with stock", "Pack an order when inventory exists.", "Inventory decreases and order moves to PACKED.", "Pass"),
        ("TC-09", "Order packing without stock", "Pack an order with insufficient inventory.", "InsufficientStockException is returned and inventory is rolled back.", "Pass"),
        ("TC-10", "Admin user management", "Create or edit an application user as ADMIN.", "User details are saved; password is hashed.", "Pass"),
        ("TC-11", "Role security", "Access admin-only endpoint as non-admin.", "Access is denied.", "Pass"),
        ("TC-12", "Frontend build", "Run npm production build.", "Build completes successfully.", "Pass"),
    ]
    add_table(doc, ["ID", "Module", "Test Scenario", "Expected Result", "Status"], rows, [0.7, 1.1, 2.0, 2.6, 0.8], font_size=7)
    doc.add_paragraph(
        "The test cases cover master data, business transactions, negative scenarios and security behavior. Negative "
        "testing is especially useful because it proves that the system does not silently accept wrong data."
    )
    uat_rows = [
        ("Warehouse Manager", "Logs in as admin, reviews inventory, manages products, suppliers and users.", "Manager can control and monitor warehouse data."),
        ("Floor Operator", "Logs in as operator, receives inbound stock and checks inventory snapshots.", "Operator can complete daily warehouse work with simple screens."),
        ("Order Processing User", "Creates customer order, starts picking, packs and ships the order.", "Order status changes correctly and stock is reduced during packing."),
        ("System Administrator", "Checks backend health endpoint and verifies PostgreSQL data in pgAdmin.", "System is running and database tables contain expected records."),
    ]
    add_table(doc, ["User", "Acceptance Activity", "Expected Acceptance Result"], uat_rows, [1.4, 2.8, 2.8], font_size=8)

    doc.add_heading("4.2 Results", level=2)
    doc.add_paragraph("The latest local verification produced the following results:")
    result_rows = [
        ("Backend tests", "mvn test", "12 tests executed, 0 failures, 0 errors."),
        ("Frontend build", "npm run build", "Production build completed successfully."),
        ("Database connection", "PostgreSQL local profile", "Application connects to PostgreSQL database wms."),
        ("Performance check", "scripts\\check-api-performance.ps1", "Measures common endpoint response time against the 200 ms target."),
    ]
    add_table(doc, ["Area", "Command/Method", "Result"], result_rows, [1.3, 2.0, 2.7], font_size=9)
    doc.add_paragraph(
        "User acceptance testing confirms that a warehouse manager can log in, manage products, receive stock, view "
        "inventory and process orders. The project also includes CI workflow configuration so tests and frontend build "
        "can be executed automatically on GitHub."
    )
    doc.add_paragraph(
        "The testing result shows that the main project flow is working. The backend tests prove that receiving stock "
        "and order rollback logic work correctly. The frontend build proves that the React application can be packaged "
        "without syntax or dependency errors."
    )
    add_labeled_paragraph(
        doc,
        "Testing limitation",
        "The current project includes important automated tests, but a final submission can be improved by adding more "
        "controller tests, security tests and browser-based frontend tests."
    )
    doc.add_page_break()


def chapter_5(doc):
    doc.add_heading("Chapter 5: Conclusion", level=1)
    doc.add_heading("5.1 Summary", level=2)
    doc.add_paragraph(
        "The Advanced Warehouse Management System successfully implements the core requirements of the Enterprise "
        "Warehouse Management System project brief. It provides a complete web application with React frontend, "
        "Spring Boot backend, PostgreSQL database, JWT security, warehouse hierarchy, receiving and putaway, inventory "
        "tracking, QR code generation, order fulfillment and testing support."
    )
    doc.add_paragraph(
        "The project improves warehouse accuracy by replacing manual stock tracking with transaction-safe database "
        "updates. The service layer protects inventory operations using transactions and locking, while role-based "
        "security keeps administrative functions restricted."
    )
    doc.add_paragraph(
        "Overall, the project demonstrates how a real business problem can be solved using a full-stack web application. "
        "It connects user-friendly screens with strong backend logic and a reliable relational database. The project "
        "also follows professional practices such as validation, exception handling, testing, Docker support and CI."
    )
    benefit_rows = [
        ("Accuracy", "Stock quantity is updated through controlled transactions instead of manual editing."),
        ("Speed", "Users can receive stock and process orders through web forms."),
        ("Visibility", "Managers can view products, bins, inventory and order status from one dashboard."),
        ("Security", "JWT login and roles protect sensitive operations."),
        ("Maintainability", "Layered code structure makes future changes easier."),
    ]
    add_table(doc, ["Benefit", "Explanation"], benefit_rows, [1.3, 4.7], font_size=9)
    add_labeled_paragraph(
        doc,
        "Learning outcome",
        "This project gives practical experience in Java backend development, React frontend development, relational "
        "database design, REST API integration, security and testing."
    )

    doc.add_heading("5.2 Future Enhancements", level=2)
    add_bullets(doc, [
        "Integrate real handheld barcode scanners or mobile camera scanning.",
        "Add advanced reporting dashboards for stock aging, fast-moving SKUs and bin utilization.",
        "Introduce Flyway or Liquibase database migrations for production-grade schema management.",
        "Add audit logs for every administrative change.",
        "Add email or SMS alerts for low stock and delayed purchase orders.",
        "Implement picking route optimization for floor operators.",
        "Add CSV/Excel import and export for product and inventory data.",
        "Add automated load testing to continuously verify the 200 ms API response target.",
        "Deploy the application to a cloud platform with environment-managed secrets.",
    ])
    doc.add_paragraph(
        "These future enhancements can make the system more useful in a real warehouse environment. The current project "
        "builds the foundation, and later improvements can focus on automation, analytics, mobile use and cloud deployment."
    )
    doc.add_page_break()


def chapter_6(doc):
    doc.add_heading("Chapter 6: References", level=1)
    refs = [
        "Infotact Technical Internship Program project specification document.",
        "Spring Boot Reference Documentation - https://spring.io/projects/spring-boot",
        "Spring Security Reference Documentation - https://spring.io/projects/spring-security",
        "Spring Data JPA Documentation - https://spring.io/projects/spring-data-jpa",
        "Hibernate ORM Documentation - https://hibernate.org/orm/documentation/",
        "PostgreSQL Documentation - https://www.postgresql.org/docs/",
        "React Documentation - https://react.dev/",
        "Vite Documentation - https://vite.dev/",
        "ZXing Barcode Library - https://github.com/zxing/zxing",
        "JUnit 5 Documentation - https://junit.org/junit5/docs/current/user-guide/",
    ]
    add_numbered(doc, refs)
    doc.add_page_break()


def chapter_7(doc):
    doc.add_heading("Chapter 7: Appendices", level=1)
    doc.add_heading("Appendix A: Run Instructions", level=2)
    doc.add_paragraph("Start backend and frontend locally:")
    add_code_block(doc, "cd \"C:\\Users\\shriyal\\OneDrive\\Documents\\New project\"\nscripts\\run-backend-local.cmd\nscripts\\run-frontend-local.cmd")
    doc.add_paragraph("Open the application:")
    add_code_block(doc, "Frontend: http://localhost:5173\nBackend health: http://localhost:8081/actuator/health")
    doc.add_paragraph(
        "The backend should be started first because the frontend depends on backend APIs for login and data. PostgreSQL "
        "must also be running before the backend starts. After both servers start, the frontend can be opened in a browser."
    )

    doc.add_heading("Appendix B: Default Development Users", level=2)
    add_table(doc, ["Username", "Role", "Purpose"], [
        ("admin", "ADMIN", "Administrator account for setup and management."),
        ("operator", "OPERATOR", "Warehouse floor operations account."),
    ], [1.2, 1.2, 3.6], font_size=9)

    doc.add_heading("Appendix C: Important API Endpoints", level=2)
    endpoints = [
        ("POST", "/api/auth/login", "Authenticate and receive JWT token."),
        ("GET/POST", "/api/warehouses", "List or create warehouses."),
        ("GET/POST", "/api/products", "List or create products."),
        ("GET", "/api/products/{id}/barcode", "Generate product QR/barcode image."),
        ("POST", "/api/inventory/receive", "Receive stock into warehouse storage."),
        ("GET", "/api/inventory/snapshots", "View current inventory position."),
        ("POST", "/api/orders", "Create customer order."),
        ("POST", "/api/orders/{id}/pack", "Pack order and decrement stock."),
        ("POST", "/api/orders/{id}/ship", "Ship packed order."),
        ("GET/POST", "/api/users", "Admin-only user management."),
    ]
    add_table(doc, ["Method", "Endpoint", "Use"], endpoints, [0.8, 2.2, 3.0], font_size=8)

    doc.add_heading("Appendix D: Project Folder Structure", level=2)
    folder_rows = [
        ("backend", "Spring Boot backend source code, tests and Maven configuration."),
        ("frontend", "React/Vite frontend source code, CSS and package configuration."),
        ("docs", "Project documentation, API request file and generated black book."),
        ("scripts", "Helper scripts for running backend, frontend and performance checks."),
        (".github/workflows", "GitHub Actions CI workflow for backend tests and frontend build."),
        ("docker-compose.yml", "Container setup for PostgreSQL, backend and frontend."),
    ]
    add_table(doc, ["Folder/File", "Purpose"], folder_rows, [1.8, 4.2], font_size=9)

    doc.add_heading("Appendix E: Security and Validation Checklist", level=2)
    checklist_rows = [
        ("JWT authentication", "Implemented", "Protected APIs require a valid token."),
        ("Role-based access", "Implemented", "ADMIN and OPERATOR roles are used."),
        ("Password hashing", "Implemented", "BCrypt is used for user passwords."),
        ("Input validation", "Implemented", "DTO validation checks required and positive fields."),
        ("CORS configuration", "Implemented", "Only configured frontend origins are allowed."),
        ("SQL injection protection", "Implemented", "JPA repositories and parameterized queries are used."),
        ("Environment secrets", "Supported", "Database and JWT values can be supplied through environment variables."),
    ]
    add_table(doc, ["Item", "Status", "Explanation"], checklist_rows, [1.5, 1.0, 3.5], font_size=8)

    doc.add_heading("Appendix F: Submission Notes", level=2)
    add_bullets(doc, [
        "Before final submission, update the index page numbers after college formatting.",
        "Replace default development secrets with environment variables before deployment.",
        "Maintain a genuine GitHub commit history across the development period.",
        "Attach screenshots of login, dashboard, product, receiving, inventory and order screens if required by the college format.",
    ])
    doc.add_page_break()


def chapter_8(doc):
    doc.add_heading("Chapter 8: Annexure - Progress Sheet", level=1)
    rows = [
        ("Week 1", "Project setup, Spring Boot scaffolding, React setup and initial database entity design.", "Completed"),
        ("Week 2", "Warehouse hierarchy, product catalog, receiving and putaway transaction logic.", "Completed"),
        ("Week 3", "QR/barcode generation, order processing, inventory decrement and exception handling.", "Completed"),
        ("Week 4", "JWT security, role-based access, frontend integration, testing, CI and documentation.", "Completed"),
    ]
    add_table(doc, ["Week", "Work Completed", "Status"], rows, [0.8, 4.3, 1.0], font_size=9)
    doc.add_paragraph(
        "Progress evidence should be supported by GitHub commits, pull requests and screenshots of the running system. "
        "The project repository includes CI configuration, Docker support, backend tests and frontend build scripts."
    )
    doc.add_paragraph(
        "The progress sheet is useful because it shows that the project was developed in stages. A staged approach is "
        "better than building everything at once. First the base project and database are created, then business logic "
        "is added, then security and frontend integration are completed, and finally testing and documentation are prepared."
    )
    detailed_rows = [
        ("Requirement Study", "Studied the WMS project brief and identified modules such as warehouse, inventory, products and orders."),
        ("Design", "Prepared layered architecture and database model based on warehouse hierarchy and ERD requirements."),
        ("Development", "Implemented backend APIs, frontend screens, PostgreSQL integration and role-based security."),
        ("Validation", "Checked forms, service logic, database records and user flows through testing."),
        ("Documentation", "Prepared black book report with project details, design, implementation, testing and future scope."),
    ]
    add_table(doc, ["Stage", "Description"], detailed_rows, [1.4, 4.6], font_size=9)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    add_cover(doc)
    add_requirement_summary(doc)
    add_index(doc)
    chapter_1(doc)
    chapter_2(doc)
    chapter_3(doc)
    chapter_4(doc)
    chapter_5(doc)
    chapter_6(doc)
    chapter_7(doc)
    chapter_8(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
